#!/usr/bin/env python3
"""Convert "Bihar_Current_Affairs_Sets1-6 (1).xlsx" into the BPSC quiz book format.

Output (bpsc/books/bihar_current_affairs/data/):
  * set_1.json … set_11.json  — one chapter file per practice set (11 sets total;
    Sets 1-6 are populated from the workbook, Sets 7-11 are empty placeholders)
  * chapters.json             — the book index consumed by bpsc/index.html

Usage:
  python3 scripts/import_bihar_ca_sets.py [--workbook PATH] [--out-dir PATH]
                                          [--total-sets 11]
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_WORKBOOK = REPO_ROOT / "Bihar_Current_Affairs_Sets1-6 (1).xlsx"
DEFAULT_OUT_DIR = REPO_ROOT / "bpsc" / "books" / "bihar_current_affairs" / "data"

SUBJECT = "Bihar Current Affairs"
EXAM = "Bihar Current Affairs Practice Set"
YEAR = "2026"
SOURCE_TYPE = "Bihar Current Affairs Sets 1–6 (Eduteria Practice Set — Excel source)"

ANSWER_RE = re.compile(r"^\(?\s*([A-Da-d])\s*[\)\.\:\-]?\s*(.*)$", re.S)

# ── topic taxonomy ────────────────────────────────────────────────────────────
# Every topic carries a list of patterns. The classifier scores a question by
# counting how many patterns match, and the highest score wins (ties are broken
# by the order below, so the more specific topics are listed first).
TOPIC_RULES = [
    ("Sports & Games", (
        r"\bhockey\b", r"\bcricket\b", r"\bkabaddi\b", r"\bfootball\b", r"\brugby\b",
        r"\bathletic", r"\bipl\b", r"\bsixes\b", r"\bolympic", r"\bkhelo\b",
        r"\bsports?\b", r"\bgrand prix\b", r"\bgold medal\b", r"\bworld cup\b",
        r"\basiad\b", r"\bnational championship\b", r"\bplayers?\b", r"\bathletes?\b",
        r"\bpara[- ]athletics\b", r"\bstadium\b", r"\bshot put\b", r"\btrophy\b",
        r"\btournament\b", r"\bkiyg\b", r"\bsports complex\b", r"\bstate games\b",
        r"\bteam\b", r"\bkushti\b", r"\bwrestl",
    )),
    ("Art, Culture & Heritage", (
        r"\bmithila\b", r"\bmadhubani\b", r"\bpainting\b", r"\bsahitya akademi\b",
        r"\bauthor\b", r"\bnovel\b", r"\bpoet", r"\bliterature\b", r"\bwriter\b",
        r"\bmemoir\b", r"\bfilm\b", r"\bculture|cultural\b",
        r"\bmanuscript\b", r"\bheritage\b", r"\bmuseum\b", r"\bshilp ?gram\b",
        r"\bcraft\b", r"\bmusic\b", r"\bdance\b", r"\btheatre\b", r"\bfolk\b",
        r"\bbuddh", r"\btemple\b", r"\bmandir\b", r"\bsanskrit\b",
        r"\btourism circuit", r"\bstatue\b", r"\bscript\b", r"\bkaithi\b",
        r"\bshilpgram", r"\bshehnai\b", r"\bhandloom\b", r"\bgi tag\b",
        r"\bunesco\b", r"\bshivling\b", r"\bghat\b", r"\bmaa |\bdevi temple\b",
        r"\bfair\b.*\bfestival\b", r"\bmehendar|\bmundeshwari|\bkundalpur|\bmandar festival",
        r"\bcivilization\b", r"\bshunga|\bkushan", r"\bsarvodaya|\bgandhi\b",
    )),
    ("Awards & Honours", (
        r"\baward(?:ed|s)?\b", r"\bhonou?red\b", r"\bhonou?r\b", r"\bprize\b",
        r"\bpadma\b", r"\bbharat ratna\b", r"\bgolden peacock\b",
        r"\bvishwakarma award\b", r"\bconferred\b", r"\bfelicitat",
        r"\bpresident'?s (?:police )?colours\b", r"\bsamman\b", r"\bgaurav\b",
        r"\bnational child award\b", r"\bquality certificate\b",
        r"\brank(?:ed|s)? (?:first|second|third|1st|2nd|3rd)\b",
    )),
    ("Law, Justice & Security", (
        r"\bhigh court\b", r"\bsupreme court\b", r"\bjudge", r"\bjudiciary\b",
        r"\bjustice\b", r"\bcourt\b", r"\bbill\b", r"\bcriminal\b",
        r"\bpolice\b", r"\bcrime", r"\bncrb\b", r"\bprison", r"\bjail\b",
        r"\bombudsman\b", r"\bcivil court", r"\bfast[- ]track court",
        r"\belection commission\b", r"\bevm\b", r"\be-?voting\b", r"\bvoters?\b",
        r"\bcivil defence\b", r"\bmock drill\b", r"\binternal security\b",
        r"\belectoral\b", r"\bprivileges committee\b", r"\bassembly elections\b",
        r"\bturnout\b", r"\bsuperintendent of police\b",
    )),
    ("Health & Sanitation", (
        r"\bhealth\b", r"\bhospital\b", r"\bdoctor", r"\bmedical\b", r"\bmedicine\b",
        r"\bsurgery|surgical\b", r"\bsurgeon\b", r"\bpatient", r"\banganwadi\b",
        r"\bvaccin", r"\bdisease\b", r"\bnutrition\b", r"\bpregnan", r"\bopd\b",
        r"\bfertility rate\b", r"\bsanitation\b", r"\btoilet\b", r"\bswachh",
        r"\bhealth card\b", r"\bcancer\b", r"\bhepatitis\b", r"\bclinic\b",
        r"\baarogya\b", r"\bigims\b", r"\baiims\b", r"\bpmch\b", r"\bstoma\b",
        r"\bgallbladder\b", r"\bcrematorium\b", r"\bpediatric\b",
    )),
    ("Environment, Forest & Wildlife", (
        r"\bclimate\b", r"\bramsar\b", r"\bgharial\b", r"\bforest", r"\bwildlife\b",
        r"\bsanctuary\b", r"\bnational park\b", r"\btiger\b", r"\bwetland",
        r"\bplantation\b", r"\bsaplings?\b", r"\bagroforestry\b", r"\bgreen cover\b",
        r"\bheat action plan\b", r"\bwaterfall", r"\bpollution\b", r"\bsewage\b",
        r"\bafforestation\b", r"\bmineral", r"\benvironment", r"\bganga\b",
        r"\bflood", r"\blion breeding\b", r"\bcrocodile\b", r"\bfish(?:eries|ery)?\b",
        r"\bvan mahotsav\b", r"\bjal jeevan hariyali\b", r"\bcarbon emission",
        r"\bwater heritage\b", r"\bcocopeat\b", r"\bnursery and green",
        r"\bwater sharing\b", r"\briver basin\b",
    )),
    ("Energy & Power", (
        r"\bpower\b", r"\belectricity\b", r"\bdiscom", r"\bsmart meter",
        r"\bsolar\b", r"\benergy\b", r"\brenewable\b", r"\bethanol\b",
        r"\bsupercomputer\b", r"\bthermal\b", r"\bsubstation\b", r"\btransmission\b",
        r"\bhydropower|hydroelectric\b", r"\bnuclear\b", r"\bcoal\b",
        r"\belectric vehicle", r"\be-vehicle", r"\bbijli\b", r"\bkwh\b",
    )),
    ("Science, IT & Innovation", (
        r"\bartificial intelligence\b", r"\bai mission\b", r"\bai impact\b",
        r"\bsemiconductor\b", r"\bdrone\b", r"\bapp\b", r"\bapps\b",
        r"\bdigital\b", r"\btechnology\b", r"\bspace club\b", r"\bsatellite\b",
        r"\brobotic\b", r"\bsoftware\b", r"\binnovation\b", r"\bstart[- ]?up",
        r"\bportal\b", r"\bisro\b", r"\bsupercomputing mission\b",
        r"\bsemen facility\b", r"\bgokul mission\b", r"\bresearch cent(?:er|re)\b",
        r"\bseed bank\b", r"\bai agent\b", r"\bai\b", r"\be-?(?:nirdesh|kshamta|sakshya)\b",
    )),
    ("Transport & Connectivity", (
        r"\bexpressway\b", r"\bmetro\b", r"\bbridge\b", r"\bhighway\b",
        r"\brailway|\brail\b", r"\bairport\b", r"\bflight\b", r"\bbullet train\b",
        r"\bcity bus\b|\bbus service\b|\bbus driver\b", r"\bwaterway\b", r"\bjetty\b",
        r"\bwater transport\b", r"\bflyover\b", r"\blogistics park\b",
        r"\bgreenfield\b", r"\btownship\b", r"\budan\b", r"\bmaritime\b",
        r"\bterminal\b", r"\bcorridor\b", r"\bcable[- ]stay", r"\bglass bridge\b",
        r"\bcommuters\b", r"\btransport department\b", r"\bpath\b",
    )),
    ("Education, Skill & Employment", (
        r"\bschool\b", r"\bschools\b", r"\beducation\b", r"\buniversity\b",
        r"\bcollege\b", r"\bstudent", r"\bskill\b", r"\bemployment\b",
        r"\bjob[s]?\b", r"\bteacher", r"\bcbse\b", r"\bcurriculum\b",
        r"\bscholarship\b", r"\bcoaching\b", r"\btraining\b", r"\bplacement\b",
        r"\bliteracy\b", r"\bhostel\b", r"\bupsc\b", r"\bcent(?:er|re) of excellence\b",
        r"\bvocational\b", r"\bkarmayogi\b", r"\bcredit card scheme\b",
        r"\binspire award\b", r"\bgram swaraj\b", r"\bemployment fair\b",
        r"\bcivil seva\b", r"\bpratigya\b", r"\byuva\b",
    )),
    ("Agriculture, Rural & Livelihoods", (
        r"\bagricultur", r"\bfarm(?:er|ing)?\b", r"\bcrop\b", r"\bcrops\b",
        r"\bmakhana\b", r"\bmillet\b", r"\bwheat\b", r"\bpaddy\b", r"\brice\b",
        r"\bseed\b", r"\bhorticultur", r"\bmango\b", r"\bturmeric\b", r"\bkisan\b",
        r"\bpanchayat\b", r"\brural\b", r"\blivestock\b", r"\bcattle\b", r"\bcows?\b",
        r"\bdairy\b", r"\birrigat", r"\bsoil\b", r"\bkrishi\b", r"\bjeevika\b",
        r"\bfood grain", r"\bsugarcane\b", r"\bsugar mill", r"\bgram panchayat\b",
        r"\bmilch\b", r"\bpublic distribution\b", r"\bfood security\b",
        r"\bants?yodaya\b", r"\bchaur region\b", r"\bagriculture roadmap\b",
    )),
    ("Industry, Investment & Infrastructure", (
        r"\bindustr", r"\bmou\b", r"\binvest(?:ment|or)", r"\bfactory\b",
        r"\bindustrial hub\b", r"\bmanufactur", r"\bmsme\b", r"\bentrepreneur",
        r"\bconclave\b", r"\bsummit\b", r"\bexport\b", r"\bbusiness\b",
        r"\bspecial economic zone|\bsez\b", r"\bdevelopment summit\b",
        r"\btrade fair\b", r"\btoy factory\b", r"\btextile\b", r"\bleather\b",
    )),
    ("Economy, Budget & Finance", (
        r"\bgdp\b", r"\bgsdp\b", r"\bbudget\b", r"\bfiscal\b", r"\brevenue\b",
        r"\bper capita\b", r"\beconom", r"\bfinance\b", r"\bbank\b", r"\btax\b",
        r"\binflation\b", r"\bgrowth rate\b", r"\bsavings\b", r"\bcrore\b",
        r"\ballocated\b", r"\bniti aayog\b", r"\bcredit[- ]deposit\b",
        r"\bmarket\b", r"\bprice\b", r"\b₹\b", r"\brupees\b", r"\bspecial assistance\b",
        r"\brating\b", r"\bexpenditure\b",
    )),
    ("Government Schemes & Welfare", (
        r"\byojana\b", r"\bscheme\b", r"\bmission\b", r"\bnis ?chay\b", r"\bnischay\b",
        r"\bpension\b", r"\bsubsidy\b", r"\bwelfare\b", r"\bassistance\b",
        r"\bbeneficiar", r"\bcabinet\b", r"\bhelpline\b", r"\blaunched\b",
        r"\binitiative\b", r"\bcampaign\b", r"\babhiyan\b", r"\bprotsahan\b",
        r"\bsamman nidhi\b", r"\bbpl\b", r"\bshivir\b", r"\bsahyog\b",
        r"\bkutir jyoti\b", r"\bsaat nischay\b", r"\bsamriddhi yatra\b",
        r"\bpragati yatra\b", r"\bmahila samvad\b", r"\bsarvodaya\b",
        r"\bcommittee\b", r"\bcommission\b", r"\bboard\b",
    )),
    ("Polity, Governance & Administration", (
        r"\bgovernor\b", r"\bchief minister\b", r"\bdeputy chief minister\b",
        r"\bminister\b", r"\bassembly\b", r"\bcouncil\b", r"\bspeaker\b",
        r"\bgovernance\b", r"\badministration\b", r"\bmunicipal\b",
        r"\bdepartment\b", r"\bsecretary\b", r"\brules\b", r"\bordinance\b",
        r"\bdistrict magistrate\b", r"\blok sabha\b", r"\brajya sabha\b",
        r"\bmla\b", r"\bcabinet formation\b", r"\bfoundation day\b",
        r"\bsession\b", r"\bbudget session\b", r"\bchief secretary\b",
        r"\bnagarpalika|\bnagar panchayat\b", r"\btranslation services\b",
    )),
    ("Social Issues, Demography & Census", (
        r"\bcensus\b", r"\bpopulation\b", r"\bpoverty\b", r"\bchild(?:ren| labour| labor)?\b",
        r"\bwomen'?s?\b", r"\bgender\b", r"\btribal\b", r"\bcaste\b", r"\bdisabled\b",
        r"\bdivyang", r"\bmissing\b", r"\bsocial\b", r"\bcaa\b", r"\bcitizenship\b",
        r"\bmigration\b|\bmigrant", r"\bunorganised\b|\bunorganized\b",
        r"\blabou?r\b", r"\bwidow\b", r"\belderly\b", r"\bsex ratio\b",
        r"\bkanya\b", r"\bbackward class", r"\bsafai karamchari\b", r"\bsurvey\b",
    )),
    ("International & National Affairs", (
        r"\bmyanmar\b", r"\bmauritius\b", r"\bsingapore\b", r"\bjapan\b", r"\bchina\b",
        r"\busa\b", r"\bgermany\b", r"\bnepal\b", r"\bbhutan\b", r"\bbangladesh\b",
        r"\bvietnam\b", r"\bafrican\b", r"\bforeign\b", r"\bbilateral\b",
        r"\bunited nations|\buno\b", r"\bindia'?s first\b", r"\bcountry'?s first\b",
        r"\brepublic day\b", r"\bprime minister\b", r"\bcentral government\b",
        r"\bsdg\b", r"\bworld'?s (?:first|largest|tallest|fifth|fourth)\b",
        r"\bparliament\b", r"\bindian-americans\b", r"\bstate status\b",
    )),
    ("Personalities, Appointments & Obituaries", (
        r"\bappointed\b", r"\bappointment\b", r"\bpassed away\b", r"\bdied\b",
        r"\bdeath\b", r"\bresident of\b", r"\bbelongs to\b", r"\bborn\b",
        r"\banniversary\b", r"\bqueen of\b", r"\btitle of\b", r"\bmrs\. india\b",
        r"\bfirst (?:person|woman|man|indian|voter)\b", r"\bwho is the author\b",
        r"\bfreedom fighter\b", r"\bsocialist thinker\b", r"\bnodal institution\b",
    )),
]

COMPILED_RULES = [(topic, [re.compile(p, re.I) for p in pats]) for topic, pats in TOPIC_RULES]
DEFAULT_TOPIC = "Miscellaneous Current Affairs"

# Manual corrections for questions where keyword scoring picks the wrong bucket.
TOPIC_OVERRIDES = {
    # (set number, question number inside the set): topic
    # ── Set 1 ────────────────────────────────────────────────────────────────
    (1, 10): "Polity, Governance & Administration",   # Amin → Assistant Revenue Officer
    (1, 13): "Science, IT & Innovation",              # village Wi-Fi by 2027
    (1, 16): "Social Issues, Demography & Census",    # unorganised workers ranking
    (1, 18): "Agriculture, Rural & Livelihoods",      # non-vegetarian food brand
    (1, 45): "Agriculture, Rural & Livelihoods",      # rare-seed seed bank
    # ── Set 2 ────────────────────────────────────────────────────────────────
    (2, 9): "Social Issues, Demography & Census",     # tribal dignity campaign
    (2, 13): "Social Issues, Demography & Census",    # Kanya Mandap in every panchayat
    (2, 17): "Health & Sanitation",                   # LaQshya quality certificate
    (2, 20): "Environment, Forest & Wildlife",        # SDG clean-water ranking
    (2, 23): "Industry, Investment & Infrastructure", # Brand Bihar initiative
    (2, 37): "Economy, Budget & Finance",             # RRB equity pattern
    (2, 44): "Education, Skill & Employment",         # Delta Ranking — education
    (2, 45): "Transport & Connectivity",              # Kacchi Dargah–Bidupur bridge
    (2, 49): "Agriculture, Rural & Livelihoods",      # Mango Centre of Excellence
    # ── Set 3 ────────────────────────────────────────────────────────────────
    (3, 2): "Energy & Power",                         # energy museum
    (3, 3): "Art, Culture & Heritage",                # Karpoori Thakur museum renamed
    (3, 9): "Industry, Investment & Infrastructure",  # AI Impact Summit MoUs
    (3, 33): "Awards & Honours",                      # Golden Peacock — Horticulture
    (3, 37): "Awards & Honours",                      # National Handloom Award
    (3, 38): "Health & Sanitation",                   # Swachhta Survey ranking
    (3, 41): "Art, Culture & Heritage",               # Bihar pavilion at trade fair
    (3, 45): "Industry, Investment & Infrastructure", # SEZ districts
    # ── Set 4 ────────────────────────────────────────────────────────────────
    (4, 10): "Polity, Governance & Administration",   # panchayat reps at R-Day parade
    (4, 15): "Law, Justice & Security",               # Bihar Police honoured
    (4, 17): "Social Issues, Demography & Census",    # Divyangjan Mahotsav
    (4, 19): "Industry, Investment & Infrastructure", # Sakri & Rayam sugar mills
    (4, 27): "Agriculture, Rural & Livelihoods",      # Bihar Saras Mela
    (4, 33): "Industry, Investment & Infrastructure", # Jamui & Nawada projects
    (4, 40): "Education, Skill & Employment",         # Inspire Award Scheme
    (4, 41): "Art, Culture & Heritage",               # state-level youth festival
    (4, 44): "Law, Justice & Security",               # fast-track courts
    (4, 46): "Science, IT & Innovation",              # Param Rudra supercomputer
    # ── Set 5 ────────────────────────────────────────────────────────────────
    (5, 15): "International & National Affairs",      # migrant workers rank (MEA)
    (5, 16): "Polity, Governance & Administration",   # first cabinet meeting
    (5, 23): "Law, Justice & Security",               # Bihar Police national honour
    (5, 35): "Social Issues, Demography & Census",    # beggar rehabilitation homes
    (5, 38): "Science, IT & Innovation",              # seismic observatory
    (5, 45): "Polity, Governance & Administration",   # Village-Settlement Chalo Abhiyan
    # ── Set 6 ────────────────────────────────────────────────────────────────
    (6, 6): "Sports & Games",                         # Bettiah sisters' medals
    (6, 12): "Polity, Governance & Administration",   # Karpoori Thakur — 1952 election
    (6, 16): "Agriculture, Rural & Livelihoods",      # Sonachur rice GI tag
    (6, 17): "Art, Culture & Heritage",               # Phulhar Sthan
    (6, 30): "Social Issues, Demography & Census",    # Safai Karamchari Commission
    (6, 36): "Agriculture, Rural & Livelihoods",      # sugarcane research centre
    (6, 42): "Agriculture, Rural & Livelihoods",      # National Makhana Board
    (6, 46): "Agriculture, Rural & Livelihoods",      # sex sorted semen facility
    (6, 49): "Agriculture, Rural & Livelihoods",      # fish production ranking
}


def classify_topic(text: str, set_no: int = 0, q_no: int = 0) -> str:
    key = (set_no, q_no)
    if key in TOPIC_OVERRIDES:
        return TOPIC_OVERRIDES[key]
    best_topic, best_score = DEFAULT_TOPIC, 0
    for topic, patterns in COMPILED_RULES:
        score = sum(1 for pattern in patterns if pattern.search(text))
        if score > best_score:
            best_topic, best_score = topic, score
    return best_topic if best_score else DEFAULT_TOPIC


STATEMENT_RE = re.compile(r"(?:^|\n)\s*(?:\(?\s*(?:[IVXivx]{1,5}|[1-4]|[a-d])\s*[\)\.\:]|\bStatement\b|\bAssertion\b)", re.M)


def classify_question_type(text: str) -> str:
    lowered = text.lower()
    if re.search(r"\bassertion\b", lowered) and re.search(r"\breason\b", lowered):
        return "Assertion and Reason — Standard Relationship Evaluation"
    if re.search(r"match list|list[- ]?i\b.*list[- ]?ii\b|match the following", lowered, re.S):
        return "Match the Following — Two-List Matching — Answer-Code Matrix"
    if re.search(r"\bcode[s]?\b|\busing the code", lowered) or STATEMENT_RE.search(text):
        if re.search(r"\bincorrect\b|\bnot correct\b|\bfalse\b|\bnot true\b", lowered):
            return "Negative/Exception MCQ — Single Correct"
        return "Multiple Statements — Correct Combination by Code"
    if re.search(r"\bincorrect\b|\bnot correct\b|\bexcept\b|\bnot true\b|\bfalse\b", lowered):
        return "Negative/Exception MCQ — Single Correct"
    return "Standard MCQ — Direct Factual/Conceptual, Single Correct"


def clean(value) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\u00a0", " ")
    return "\n".join(line.strip() for line in text.split("\n")).strip()


def parse_answer(raw: str):
    """Return (letter, answer_text) from strings like 'B) Multi Post S-3'."""
    text = clean(raw)
    if not text:
        return "", ""
    match = ANSWER_RE.match(text)
    if not match:
        return "", text
    return match.group(1).upper(), match.group(2).strip()


def read_workbook(path: Path):
    try:
        import openpyxl
    except ImportError:  # pragma: no cover
        sys.exit("openpyxl is required: pip install openpyxl")
    workbook = openpyxl.load_workbook(path, data_only=True)
    sheets = []
    for ws in workbook.worksheets:
        match = re.search(r"(\d+)", ws.title)
        set_no = int(match.group(1)) if match else len(sheets) + 1
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[2] is None and row[1] is None:
                continue
            rows.append(row)
        sheets.append((set_no, ws.title, rows))
    sheets.sort(key=lambda item: item[0])
    return sheets


def build_question(row, set_no: int, seq: int, duplicate_of: int = 0) -> dict:
    _set_no, q_no, stem, opt_a, opt_b, opt_c, opt_d, raw_answer = row[:8]
    question_text = clean(stem)
    options = {
        "A": clean(opt_a),
        "B": clean(opt_b),
        "C": clean(opt_c),
        "D": clean(opt_d),
    }
    answer, answer_text = parse_answer(raw_answer)
    if answer not in options:
        answer = ""
    q_no = int(q_no) if isinstance(q_no, (int, float)) else seq
    return {
        "q": question_text,
        "question_type": classify_question_type(question_text),
        "options": options,
        "answer": answer,
        "explanation": "",
        "exam": EXAM,
        "year": YEAR,
        "answer_key_text": f"{answer}: {answer_text}" if answer_text else "",
        "note": (
            f"Source: Bihar Current Affairs Set {set_no}, Q{q_no} — answer key as published in "
            f"the Eduteria practice workbook"
            + (f". Repeated verbatim in the source workbook — also Q{duplicate_of} of this set."
               if duplicate_of else ".")
        ),
        "original_id": (set_no - 1) * 50 + q_no,
        "set_number": set_no,
        "set_question_no": q_no,
        "source_type": SOURCE_TYPE,
        "source_question_type": "",
        "difficulty": "Medium",
        "topic": classify_topic(question_text, set_no, q_no),
        "keyFact": "",
        "trap": "",
        "duplicate_of": duplicate_of,
        "tag": f"Bihar Current Affairs — Set {set_no}",
        "section": f"Set {set_no}",
    }


def norm_text(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", text.lower())


def find_duplicates(rows) -> dict:
    """Map 0-based row index → 1-based question number of the first identical stem."""
    seen, dupes = {}, {}
    for i, row in enumerate(rows):
        key = norm_text(clean(row[2]))
        if not key:
            continue
        if key in seen:
            dupes[i] = seen[key] + 1
        else:
            seen[key] = i
    return dupes


def write_json(path: Path, payload) -> None:
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--workbook", default=str(DEFAULT_WORKBOOK))
    parser.add_argument("--out-dir", default=str(DEFAULT_OUT_DIR))
    parser.add_argument("--total-sets", type=int, default=11, help="Number of sets the book exposes (default: 11)")
    args = parser.parse_args()

    workbook = Path(args.workbook)
    out_dir = Path(args.out_dir)
    if not workbook.is_file():
        sys.exit(f"workbook not found: {workbook}")
    out_dir.mkdir(parents=True, exist_ok=True)

    sheets = read_workbook(workbook)
    index = []
    total_questions = 0

    populated = {set_no: rows for set_no, _title, rows in sheets}

    for set_no in range(1, args.total_sets + 1):
        rows = populated.get(set_no, [])
        dupes = find_duplicates(rows)
        questions = [
            build_question(row, set_no, i + 1, dupes[i] if i in dupes else 0)
            for i, row in enumerate(rows)
        ]
        for i in sorted(dupes):
            print(f"  ! Set {set_no}: Q{i + 1} repeats Q{dupes[i]} (kept as in the source workbook)")
        total_questions += len(questions)
        file_name = f"set_{set_no}.json"
        write_json(out_dir / file_name, {
            "subject": f"{SUBJECT} — Set {set_no}",
            "set_number": set_no,
            "chapters": [{
                "chapter_name": f"Set {set_no}",
                "questions": questions,
            }],
        })
        index.append({
            "key": f"set_{set_no}",
            "label": f"Set {set_no}",
            "emoji": "📰",
            "file": file_name,
            "chapters": [{"name": f"Set {set_no}", "count": len(questions)}],
        })

    write_json(out_dir / "chapters.json", index)

    print(f"wrote {len(index)} sets → {out_dir}")
    for set_no, title, rows in sheets:
        print(f"  Set {set_no} ({title}): {len(rows)} questions")
    print(f"total questions: {total_questions}")

    # quick sanity report
    topic_counts = {}
    type_counts = {}
    for set_no, _title, rows in sheets:
        for i, row in enumerate(rows):
            q = build_question(row, set_no, i + 1)
            topic_counts[q["topic"]] = topic_counts.get(q["topic"], 0) + 1
            type_counts[q["question_type"]] = type_counts.get(q["question_type"], 0) + 1
    print("\ntopics:")
    for topic, count in sorted(topic_counts.items(), key=lambda kv: -kv[1]):
        print(f"  {count:>4}  {topic}")
    print("question types:")
    for qt, count in sorted(type_counts.items(), key=lambda kv: -kv[1]):
        print(f"  {count:>4}  {qt}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
